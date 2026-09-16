"""internal_panel.xlsx → 표준 스키마.

확인된 해석 (헤더가 데이터보다 한 칸 왼쪽으로 밀려 있습니다)

    A열  권역        병합셀. 표준 스키마에 자리가 없어 버립니다
    B열  시도명      references/region-codes.csv 로 코드 변환
    C·D·E  합계출산율  2023 · 2024 · 2025
    F·G·H  출생아수    2023 · 2024 · 2025

    r1 제목 · r2 빈행 · r3~r4 헤더 · r5~r14 데이터 · r15 빈행 · r16~r18 각주

규칙
- 출생아수는 명으로 통일합니다. '43.5천' 같은 천 단위 표기는 1000 을 곱합니다(주1).
- '-' 는 NA_CONFIDENTIAL, 공란은 NA_NOTSURVEYED 입니다(주2).
- 2025년은 vintage 잠정, 그 외는 확정입니다(주3).
- 지역명이 사전에 없으면 임의로 매핑하지 않고 멈춥니다.
- 각주 행은 데이터로 읽지 않습니다.

    python exercises/03/convert_panel.py
"""
from __future__ import annotations

import os
import re
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, HERE)
sys.path.insert(0, BASE)

from contract import COLUMNS, Row, SchemaError  # noqa: E402
from scripts.normalize import load_region_map  # noqa: E402

XLSX = os.path.join(HERE, "internal_panel.xlsx")
SHEET = "시도별 출생 현황"
SOURCE = "INTERNAL"
UNIT = "명"

FIRST_DATA_ROW = 5
REGION_COL = 2                       # B열 시도명
GROUP_COL = 1                        # A열 권역. 버리지만 보고는 합니다
FOOTNOTE = re.compile(r"^주\d*\)")
THOUSANDS = re.compile(r"^([\d.]+)\s*천$")   # '43.5천'

# (지표코드, {엑셀 열번호: 연도})
BLOCKS = [
    ("TFR", {3: 2023, 4: 2024, 5: 2025}),      # C·D·E
    ("BIRTHS", {6: 2023, 7: 2024, 8: 2025}),   # F·G·H
]
MISSING_MARK = {"-": "NA_CONFIDENTIAL"}        # 공란은 셀이 None 으로 옵니다


class ConvertError(RuntimeError):
    """해석할 수 없는 셀. 추측하지 않고 멈춥니다."""


def parse_cell(raw, indicator, coord):
    """셀 하나 → (값, 결측사유, 환산메모). 규칙 밖이면 멈춥니다."""
    if raw is None:                                   # 공란 = 미집계(주2)
        return None, "NA_NOTSURVEYED", None

    if isinstance(raw, str):
        s = raw.strip()
        if s in MISSING_MARK:                         # '-' = 비공개(주2)
            return None, MISSING_MARK[s], None
        if not s:
            return None, "NA_NOTSURVEYED", None
        m = THOUSANDS.match(s)
        if m:
            if indicator != "BIRTHS":                 # 합계출산율에 천 단위는 말이 안 됩니다
                raise ConvertError(f"{coord}: {raw!r} — {indicator} 열에 천 단위 표기가 있습니다")
            return float(m.group(1)) * 1000, None, f"{coord} {s} → {float(m.group(1)) * 1000:.0f}"
        raise ConvertError(
            f"{coord}: {raw!r} 를 해석할 수 없습니다. "
            "각주에 없는 표기이므로 추측하지 않고 멈춥니다."
        )

    if isinstance(raw, bool) or not isinstance(raw, (int, float)):
        raise ConvertError(f"{coord}: {raw!r} <{type(raw).__name__}> 는 수치가 아닙니다")
    return float(raw), None, None


def convert(log=print):
    region_map = load_region_map(BASE)
    ws = load_workbook(XLSX, data_only=True)[SHEET]

    rows, converted, dropped_groups, footnotes = [], [], [], []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        group = ws.cell(row=r, column=GROUP_COL).value
        name = ws.cell(row=r, column=REGION_COL).value

        # 각주 행은 데이터로 읽지 않습니다. A열에 문장이 들어옵니다.
        if isinstance(group, str) and FOOTNOTE.match(group.strip()):
            footnotes.append(f"r{r}: {group.strip()}")
            continue
        if name is None:                              # 데이터 블록 끝(r15 빈 행)
            continue

        code = region_map.get(str(name).strip())
        if code is None:                              # 임의 매핑 금지
            raise ConvertError(
                f"B{r}: '{name}' 이 references/region-codes.csv 에 없습니다. "
                "임의로 매핑하지 않고 멈춥니다. 사전에 등록한 뒤 다시 실행하세요."
            )
        if isinstance(group, str) and group.strip():
            dropped_groups.append(f"r{r} {group.strip()} → {name}")

        for indicator, year_by_col in BLOCKS:
            for col, year in year_by_col.items():
                coord = f"{get_column_letter(col)}{r}"
                value, reason, memo = parse_cell(
                    ws.cell(row=r, column=col).value, indicator, coord)
                if memo:
                    converted.append(f"{name} {year} {memo}")
                rows.append(Row(
                    source=SOURCE,
                    indicator_code=indicator,
                    region_code=code,
                    period=year,
                    value=value,
                    unit=UNIT,
                    vintage="잠정" if year == 2025 else "확정",   # 주3
                    source_url=os.path.relpath(XLSX, BASE).replace("\\", "/"),
                    missing_reason=reason,
                ))

    if footnotes:
        log("각주 행 제외:")
        for f in footnotes:
            log(f"  {f}")
    return rows, converted, dropped_groups


def main():
    try:
        rows, converted, dropped = convert()
    except (ConvertError, SchemaError) as e:
        print(f"중단: {e}", file=sys.stderr)
        return 1

    df = pd.DataFrame([r.as_dict() for r in rows], columns=COLUMNS)
    out = os.path.join(BASE, "out", "internal_panel.csv")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    df.to_csv(out, index=False, encoding="utf-8-sig")

    print()
    print(f"{len(df)}행 · 시도 {df['region_code'].nunique()}개 · 지표 "
          f"{sorted(df['indicator_code'].unique())} · 기간 "
          f"{df['period'].min()}~{df['period'].max()} → out/internal_panel.csv")
    print()
    print("단위 환산 (천명 → 명):")
    for c in converted:
        print(f"  {c}")
    print()
    print("결측:")
    miss = df[df["value"].isna()]
    print(miss[["indicator_code", "region_code", "period", "missing_reason"]].to_string(index=False))
    print()
    print("vintage:")
    print(df.groupby(["period", "vintage"]).size().to_string())
    print()
    print(f"버린 A열 권역 {len(dropped)}건 (표준 스키마에 자리 없음):")
    for d in dropped:
        print(f"  {d}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
